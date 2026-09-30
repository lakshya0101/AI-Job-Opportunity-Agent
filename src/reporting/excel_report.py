from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional, Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.collectors.base import CollectorStatus, SourceHealth
from src.models.job import Job
from src.reporting.report_builder import DailyReport

# Style definitions
HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

SECTION_FILL = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

SUCCESS_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
SUCCESS_FONT = Font(name="Calibri", size=11, bold=True, color="166534")

FAILED_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
FAILED_FONT = Font(name="Calibri", size=11, bold=True, color="991B1B")

NOT_IMPL_FILL = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
NOT_IMPL_FONT = Font(name="Calibri", size=11, color="6B7280", italic=True)

LINK_FONT = Font(name="Calibri", size=11, color="1D4ED8", underline="single")
REGULAR_FONT = Font(name="Calibri", size=11, color="0F172A")
BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

JOB_COLUMNS = [
    ("Company", 22),
    ("Role", 28),
    ("Location", 20),
    ("Work Mode", 14),
    ("Experience", 16),
    ("Eligibility", 16),
    ("Compensation", 18),
    ("Posting Date", 14),
    ("Deadline", 14),
    ("Match Score", 13),
    ("Role Score", 12),
    ("Skill Score", 12),
    ("Loc Score", 12),
    ("Exp Score", 12),
    ("Fresh Score", 12),
    ("Matched Skills", 35),
    ("Match Reason", 35),
    ("Source", 22),
    ("Application Link", 20),
    ("Careers Link", 20),
    ("Recruiter Name", 18),
    ("Recruiter Email", 22),
    ("New", 8),
    ("Updated", 10),
    ("Urgent", 8),
]


def _format_cell(cell, font=REGULAR_FONT, fill=None, alignment=None, border=THIN_BORDER):
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border


def _auto_fit_columns(ws, max_width_limit: int = 50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), max_width_limit)


def _build_summary_sheet(ws, report: DailyReport, report_date: str):
    ws.title = "Summary"
    ws.views.sheetView[0].showGridLines = True

    # Title
    ws["A1"] = "AI JOB OPPORTUNITY REPORT"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    
    ws["A2"] = f"Report Date: {report_date}"
    ws["A2"].font = Font(name="Calibri", size=11, italic=True, color="64748B")

    ws["A3"] = f"Generated UTC: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="94A3B8")

    # Metrics Section
    ws["A5"] = "PIPELINE SUMMARY"
    ws.merge_cells("A5:B5")
    _format_cell(ws["A5"], font=SECTION_FONT, fill=SECTION_FILL, alignment=Alignment(horizontal="center"))
    _format_cell(ws["B5"], fill=SECTION_FILL)

    metrics = [
        ("Total Jobs Collected", report.summary.collected),
        ("Total Normalized", report.summary.normalized),
        ("Unique Jobs (Deduplicated)", report.summary.deduplicated),
        ("Newly Discovered (Since Last Run)", report.summary.new),
        ("Meaningfully Updated", report.summary.updated),
        ("Matched User Criteria", report.summary.matched),
        ("Final Reportable Opportunities", report.total_jobs),
    ]

    for idx, (label, val) in enumerate(metrics, start=6):
        ws[f"A{idx}"] = label
        ws[f"B{idx}"] = val
        _format_cell(ws[f"A{idx}"], font=BOLD_FONT, alignment=Alignment(horizontal="left"))
        _format_cell(ws[f"B{idx}"], font=BOLD_FONT, alignment=Alignment(horizontal="center"))

    # Sections Breakdown Table
    ws["D5"] = "REPORT SECTIONS BREAKDOWN"
    ws.merge_cells("D5:E5")
    _format_cell(ws["D5"], font=SECTION_FONT, fill=SECTION_FILL, alignment=Alignment(horizontal="center"))
    _format_cell(ws["E5"], fill=SECTION_FILL)

    breakdown = [
        ("Apply First (Top Matches)", len(report.apply_first)),
        ("New Today (Posted Today)", len(report.new_today)),
        ("Meaningfully Updated", len(report.updated)),
        ("Urgent (Deadline <= 48h)", len(report.urgent)),
        ("Fresh & Active (1-7 Days)", len(report.fresh_active)),
        ("Remote India Opportunities", len(report.remote_india)),
        ("Business / BI Analyst", len(report.business_analyst)),
        ("Other Strong Matches (Score >= 80)", len(report.other_strong_matches)),
    ]

    for idx, (label, val) in enumerate(breakdown, start=6):
        ws[f"D{idx}"] = label
        ws[f"E{idx}"] = val
        _format_cell(ws[f"D{idx}"], font=REGULAR_FONT, alignment=Alignment(horizontal="left"))
        _format_cell(ws[f"E{idx}"], font=BOLD_FONT, alignment=Alignment(horizontal="center"))

    # Candidate Profile Target Summary
    ws["A15"] = "CANDIDATE TARGET CRITERIA"
    ws.merge_cells("A15:E15")
    _format_cell(ws["A15"], font=HEADER_FONT, fill=HEADER_FILL, alignment=Alignment(horizontal="left"))
    for col in ["B", "C", "D", "E"]:
        _format_cell(ws[f"{col}15"], fill=HEADER_FILL)

    profile_info = [
        ("Degree & Experience", "B.Tech CSE (Data Science) | Fresher / Trainee / 0–2 Years Experience"),
        ("Target Roles", "AI/ML Engineer, GenAI/LLM Engineer, Python Backend, Data Scientist, BI/Data Analyst"),
        ("Target Locations", "Noida, New Delhi, Gurugram, Bangalore, Jaipur, Pune, Mumbai, Remote India"),
    ]

    for idx, (label, val) in enumerate(profile_info, start=16):
        ws[f"A{idx}"] = label
        ws[f"B{idx}"] = val
        ws.merge_cells(f"B{idx}:E{idx}")
        _format_cell(ws[f"A{idx}"], font=BOLD_FONT)
        _format_cell(ws[f"B{idx}"], font=REGULAR_FONT)
        for c in ["C", "D", "E"]:
            _format_cell(ws[f"{c}{idx}"])

    _auto_fit_columns(ws, max_width_limit=60)


def _build_source_health_sheet(ws, source_health: List[SourceHealth]):
    ws.title = "Source Health"
    ws.views.sheetView[0].showGridLines = True

    headers = ["Source Name", "Execution Status", "Jobs Collected", "Duration (Seconds)", "Error Message"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        _format_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, alignment=Alignment(horizontal="center"))

    for row_idx, sh in enumerate(source_health, start=2):
        status_str = sh.status.value if isinstance(sh.status, CollectorStatus) else str(sh.status)

        c1 = ws.cell(row=row_idx, column=1, value=sh.source)
        c2 = ws.cell(row=row_idx, column=2, value=status_str)
        c3 = ws.cell(row=row_idx, column=3, value=sh.job_count)
        c4 = ws.cell(row=row_idx, column=4, value=round(sh.duration_seconds, 2) if sh.duration_seconds > 0 else "—")
        c5 = ws.cell(row=row_idx, column=5, value=sh.error or "")

        # Format status badge
        if status_str == "SUCCESS":
            status_fill, status_font = SUCCESS_FILL, SUCCESS_FONT
        elif status_str == "FAILED":
            status_fill, status_font = FAILED_FILL, FAILED_FONT
        elif status_str == "NOT_IMPLEMENTED":
            status_fill, status_font = NOT_IMPL_FILL, NOT_IMPL_FONT
        else:
            status_fill, status_font = None, REGULAR_FONT

        _format_cell(c1, font=BOLD_FONT)
        _format_cell(c2, font=status_font, fill=status_fill, alignment=Alignment(horizontal="center"))
        _format_cell(c3, font=REGULAR_FONT, alignment=Alignment(horizontal="center"))
        _format_cell(c4, font=REGULAR_FONT, alignment=Alignment(horizontal="center"))
        _format_cell(c5, font=Font(name="Calibri", size=10, color="DC2626"))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _auto_fit_columns(ws, max_width_limit=50)


def _build_job_sheet(ws, sheet_title: str, jobs: List[Job]):
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True

    # Header row
    for col_idx, (col_name, _) in enumerate(JOB_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        _format_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, alignment=Alignment(horizontal="center"))

    if not jobs:
        empty_cell = ws.cell(row=2, column=1, value="No matching jobs in this category.")
        _format_cell(empty_cell, font=Font(name="Calibri", size=11, italic=True, color="64748B"))
        ws.freeze_panes = "A2"
        _auto_fit_columns(ws)
        return

    # Data rows
    for row_idx, job in enumerate(jobs, start=2):
        skills_str = ", ".join(job.skills) if job.skills else ""
        
        # Determine application hyperlink
        app_url = job.application_url or ""
        careers_url = job.careers_url or ""

        row_data = [
            job.company or "",
            job.title or "",
            job.location or "",
            job.work_mode or "",
            job.experience or "",
            job.eligibility or "",
            job.compensation or "",
            job.posting_date or "",
            job.deadline or "",
            job.match_score,
            job.role_score,
            job.skill_score,
            job.location_score,
            job.experience_score,
            job.freshness_score,
            skills_str,
            job.match_reason or "",
            job.source or "",
            app_url,
            careers_url,
            job.recruiter_name or "",
            job.recruiter_email or "",
            "Yes" if job.is_new else "No",
            "Yes" if job.is_updated else "No",
            "Yes" if job.is_urgent else "No",
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            col_name = JOB_COLUMNS[col_idx - 1][0]

            # Hyperlink handling for Application Link and Careers Link
            if col_name == "Application Link":
                if app_url:
                    cell.hyperlink = app_url
                    cell.value = "Apply Link"
                    _format_cell(cell, font=LINK_FONT, alignment=Alignment(horizontal="center"))
                elif careers_url:
                    cell.hyperlink = careers_url
                    cell.value = "Careers Page"
                    _format_cell(cell, font=LINK_FONT, alignment=Alignment(horizontal="center"))
                else:
                    cell.value = ""
                    _format_cell(cell, alignment=Alignment(horizontal="center"))
            elif col_name == "Careers Link":
                if careers_url:
                    cell.hyperlink = careers_url
                    cell.value = "Careers Page"
                    _format_cell(cell, font=LINK_FONT, alignment=Alignment(horizontal="center"))
                else:
                    cell.value = ""
                    _format_cell(cell, alignment=Alignment(horizontal="center"))
            elif col_name in {"Match Score", "Role Score", "Skill Score", "Loc Score", "Exp Score", "Fresh Score"}:
                _format_cell(cell, font=BOLD_FONT if col_name == "Match Score" else REGULAR_FONT, alignment=Alignment(horizontal="center"))
            elif col_name in {"Posting Date", "Deadline", "New", "Updated", "Urgent"}:
                _format_cell(cell, alignment=Alignment(horizontal="center"))
            elif col_name in {"Matched Skills", "Match Reason"}:
                _format_cell(cell, alignment=Alignment(wrap_text=True))
            else:
                _format_cell(cell)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    _auto_fit_columns(ws, max_width_limit=45)


def get_default_excel_filename(report_date: Optional[str] = None) -> str:
    """Generate a sanitized cross-platform filename for the Excel report."""
    if report_date:
        # Sanitize any special characters (like slashes or colons)
        date_clean = "".join(c if c.isalnum() or c in "-_" else "_" for c in report_date.strip())
        return f"ai_job_opportunities_{date_clean}.xlsx"
    
    current_date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    return f"ai_job_opportunities_{current_date}.xlsx"


def generate_excel_report(
    report: DailyReport,
    output_path: Union[str, Path],
    report_date: Optional[str] = None,
) -> Path:
    """
    Generate a formatted multi-sheet Excel workbook from DailyReport data.
    
    Sheets created:
    1. Summary
    2. Apply First
    3. New Today (if jobs present)
    4. Updated (if jobs present)
    5. Urgent (if jobs present)
    6. Fresh & Active (if jobs present)
    7. Remote India (if jobs present)
    8. Business Analyst (if jobs present)
    9. Other Strong Matches (if jobs present)
    10. Source Health
    """
    date_str = report_date or datetime.now(timezone.utc).strftime("%d %b %Y")
    path_obj = Path(output_path).resolve()
    path_obj.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # First default sheet
    ws_summary = wb.active
    _build_summary_sheet(ws_summary, report, date_str)

    # Job category sheets
    job_sections = [
        ("Apply First", report.apply_first),
        ("New Today", report.new_today),
        ("Updated", report.updated),
        ("Urgent", report.urgent),
        ("Fresh & Active", report.fresh_active),
        ("Remote India", report.remote_india),
        ("Business Analyst", report.business_analyst),
        ("Other Strong Matches", report.other_strong_matches),
    ]

    for title, job_list in job_sections:
        # Create sheet if jobs exist or if it's Apply First (which should always be visible)
        if job_list or title == "Apply First":
            ws_job = wb.create_sheet(title=title)
            _build_job_sheet(ws_job, title, job_list)

    # Source Health sheet
    ws_health = wb.create_sheet(title="Source Health")
    _build_source_health_sheet(ws_health, report.source_health)

    wb.save(path_obj)
    return path_obj
