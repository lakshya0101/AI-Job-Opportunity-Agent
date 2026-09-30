from pathlib import Path
import openpyxl

from src.collectors.base import CollectorStatus, SourceHealth
from src.models.job import Job
from src.pipeline import PipelineResult
from src.reporting.excel_report import generate_excel_report, get_default_excel_filename
from src.reporting.report_builder import build_report


def test_excel_report_creation_and_open(tmp_path):
    """Verify Excel workbook is created and can be opened by openpyxl."""
    output_file = tmp_path / "test_report.xlsx"

    job1 = Job(
        company="OpenAI",
        title="AI Engineer",
        location="Remote India",
        work_mode="Remote",
        experience="0-2 years",
        compensation="35 LPA",
        posting_date="2026-10-01",
        match_score=95.0,
        role_score=30.0,
        skill_score=25.0,
        location_score=20.0,
        experience_score=10.0,
        freshness_score=10.0,
        skills=["Python", "LLM", "LangChain"],
        application_url="https://openai.com/apply/1",
        is_new=True,
    )
    job2 = Job(
        company="Microsoft",
        title="Business Analyst",
        location="Noida",
        match_score=82.0,
        careers_url="https://microsoft.com/careers",
        is_updated=True,
    )
    job3 = Job(
        company="Local Startup",
        title="Python Dev",
        location="Gurugram",
        match_score=70.0,
    )

    source_health = [
        SourceHealth(source="official_careers", status=CollectorStatus.SUCCESS, job_count=10, duration_seconds=1.5),
        SourceHealth(source="LinkedIn", status=CollectorStatus.NOT_IMPLEMENTED, job_count=0),
        SourceHealth(source="FailingSource", status=CollectorStatus.FAILED, job_count=0, error="Connection timeout"),
    ]

    report = build_report([job1, job2, job3], source_health=source_health)
    path = generate_excel_report(report, output_file, report_date="01 Oct 2026")

    assert path.exists()

    # Open with openpyxl and inspect sheets
    wb = openpyxl.load_workbook(path)
    sheet_names = wb.sheetnames

    assert "Summary" in sheet_names
    assert "Source Health" in sheet_names
    assert "Apply First" in sheet_names
    assert "Remote India" in sheet_names
    assert "Business Analyst" in sheet_names


def test_excel_summary_sheet_content(tmp_path):
    """Verify Summary sheet contains correct metrics."""
    output_file = tmp_path / "summary_test.xlsx"

    pipe_res = PipelineResult(
        collected_count=50,
        normalized_count=48,
        deduplicated_count=40,
        new_count=12,
        updated_count=4,
        matched_count=20,
        filtered_count=8,
        source_health=[],
        reportable_jobs=[],
    )

    report = build_report([], pipeline_result=pipe_res)
    generate_excel_report(report, output_file, report_date="01 Oct 2026")

    wb = openpyxl.load_workbook(output_file)
    ws = wb["Summary"]

    assert ws["A1"].value == "AI JOB OPPORTUNITY REPORT"
    assert "01 Oct 2026" in str(ws["A2"].value)

    # Check metric values in column B
    metric_values = [ws[f"B{row}"].value for row in range(6, 13)]
    assert 50 in metric_values  # collected
    assert 40 in metric_values  # deduplicated
    assert 12 in metric_values  # new
    assert 4 in metric_values   # updated


def test_excel_source_health_sheet(tmp_path):
    """Verify Source Health sheet retains statuses, counts, durations, and errors."""
    output_file = tmp_path / "health_test.xlsx"

    source_health = [
        SourceHealth(source="Lever", status=CollectorStatus.SUCCESS, job_count=25, duration_seconds=2.1),
        SourceHealth(source="Naukri", status=CollectorStatus.NOT_IMPLEMENTED, job_count=0),
        SourceHealth(source="ScraperX", status=CollectorStatus.FAILED, job_count=0, error="HTTP 403 Forbidden"),
    ]

    report = build_report([], source_health=source_health)
    generate_excel_report(report, output_file)

    wb = openpyxl.load_workbook(output_file)
    ws = wb["Source Health"]

    assert ws.cell(row=2, column=1).value == "Lever"
    assert ws.cell(row=2, column=2).value == "SUCCESS"
    assert ws.cell(row=2, column=3).value == 25

    assert ws.cell(row=3, column=1).value == "Naukri"
    assert ws.cell(row=3, column=2).value == "NOT_IMPLEMENTED"

    assert ws.cell(row=4, column=1).value == "ScraperX"
    assert ws.cell(row=4, column=2).value == "FAILED"
    assert "HTTP 403" in str(ws.cell(row=4, column=5).value)


def test_excel_hyperlinks_and_fallback(tmp_path):
    """Verify application URL hyperlink, careers URL fallback, and blank handling."""
    output_file = tmp_path / "links_test.xlsx"

    job_direct = Job(company="A", title="Role A", location="Noida", application_url="https://direct.com/job1")
    job_careers = Job(company="B", title="Role B", location="Noida", careers_url="https://company.com/careers")
    job_none = Job(company="C", title="Role C", location="Noida")

    report = build_report([job_direct, job_careers, job_none])
    generate_excel_report(report, output_file)

    wb = openpyxl.load_workbook(output_file)
    ws = wb["Apply First"]

    # Row 2: job_direct
    cell_app1 = ws.cell(row=2, column=19)  # Application Link column
    assert cell_app1.hyperlink is not None
    assert cell_app1.hyperlink.target == "https://direct.com/job1"

    # Row 3: job_careers (falls back to careers page in Application Link)
    cell_app2 = ws.cell(row=3, column=19)
    assert cell_app2.hyperlink is not None
    assert cell_app2.hyperlink.target == "https://company.com/careers"

    # Row 4: job_none
    cell_app3 = ws.cell(row=4, column=19)
    assert cell_app3.hyperlink is None
    assert cell_app3.value in (None, "")



def test_empty_report_excel_generation(tmp_path):
    """Empty report must generate successfully with Summary, Apply First, and Source Health."""
    output_file = tmp_path / "empty_report.xlsx"

    report = build_report([])
    path = generate_excel_report(report, output_file)

    assert path.exists()
    wb = openpyxl.load_workbook(path)
    assert "Summary" in wb.sheetnames
    assert "Source Health" in wb.sheetnames
    assert "Apply First" in wb.sheetnames

    ws_apply = wb["Apply First"]
    assert "No matching jobs" in str(ws_apply.cell(row=2, column=1).value)


def test_default_filename_sanitization():
    """Verify default filename generator creates sanitized cross-platform filenames."""
    fn1 = get_default_excel_filename("01/10/2026")
    assert "/" not in fn1
    assert fn1.endswith(".xlsx")

    fn2 = get_default_excel_filename()
    assert fn2.startswith("ai_job_opportunities_")
    assert fn2.endswith(".xlsx")
